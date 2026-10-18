#!/usr/bin/env python3
"""
Full-text search for companion applications (design purpose, Step 3; Supplementary Notes S5.2.3).

This script searches the full texts of all articles evaluating client-facing applications
for the companion application names and generic terms in data/auxiliary/companion_screen_terms.yaml.
Every hit is written to a CSV file (study, term, page, context) for verification by two reviewers.

Optionally (--highlight-dir), PDFs with at least one hit are copied to a folder with the hits
highlighted (yellow = application name, green = generic term), together with a coding sheet for
the reviewers. The coding sheet is only created if it does not exist yet, so entered codes are kept.
The completed coding sheet is data/unprocessed/companion_screen_coding.xlsx, which is merged into
final_data.csv in data cleaning step 10.

Note: the full-text PDFs are not distributed with this repository for copyright reasons, so this
script can only be re-run with a local copy of the PDFs. Its output is included as
data/intermediate/companion_screen_hits.csv.

Usage:
    python 1_search_companion_terms.py --pdf-dir DIR --pdf-index FILE [--highlight-dir DIR]
"""

import argparse
import os
import re
import sys
from pathlib import Path

import pandas as pd
import pymupdf
import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

YELLOW, GREEN = (1, 1, 0), (0.5, 1, 0.5)
CODING_SHEET = "companion_screen_coding.xlsx"


def load_terms(path: Path) -> list:
    """Load the search terms as a list of (term, compiled pattern, term type)."""
    config = yaml.safe_load(open(path))
    terms = []
    for section, term_type in [("application_names", "name"), ("generic_terms", "generic")]:
        for term, spec in config[section].items():
            flags = 0 if spec.get("case_sensitive", False) else re.IGNORECASE
            terms.append((term, re.compile(spec["pattern"], flags), term_type))
    return terms


def search_pdf(pdf_path: Path, study: dict, terms: list, highlight_path: Path = None) -> list:
    """Search one PDF page by page and return its hits; optionally save a highlighted copy."""
    doc = pymupdf.open(pdf_path)
    hits = []
    for page in doc:
        # Join the words of the page into one string, remembering where each word starts and ends
        words = page.get_text("words")
        text, spans = "", []
        for w in words:
            spans.append((len(text), len(text) + len(w[4])))
            text += w[4] + " "

        for term, pattern, term_type in terms:
            for m in pattern.finditer(text):
                if highlight_path:
                    rects = [pymupdf.Rect(words[i][:4]) for i, (a, b) in enumerate(spans) if a < m.end() and b > m.start()]
                    annot = page.add_highlight_annot(rects)
                    annot.set_colors(stroke=YELLOW if term_type == "name" else GREEN)
                    annot.set_info(title="Companion screen", content=f"{term} ({term_type})")  # shown in the PDF reader's comment list
                    annot.update()
                hits.append({
                    "study_id": study["study_id"],
                    "first_author": study["first_author"],
                    "year": study["year"],
                    "term_type": term_type,
                    "term": term,
                    "page": page.number + 1,
                    "n_pages": len(doc),  # hits on the last pages are likely in the reference list
                    "context": text[max(0, m.start() - 150):m.end() + 150],
                })
    if hits and highlight_path:
        doc.save(highlight_path)
    doc.close()
    return hits


def create_coding_sheet(hits: pd.DataFrame, pdf_files: dict, path: Path) -> None:
    """Create the coding sheet for the two reviewers (one row per article with at least one hit)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Coding"
    ws.append(["No.", "First author, year", "Study ID", "PDF file", "Application names found (pages)",
               "Generic terms found (pages)", "Evaluates AI companion application – Reviewer 1",
               "Evaluates AI companion application – Reviewer 2", "Consensus", "Notes"])

    for study_id in sorted(hits["study_id"].unique(), key=lambda i: pdf_files[i]):  # same order as the PDFs in the folder
        study_hits = hits[hits["study_id"] == study_id]
        found = {}
        for term_type in ["name", "generic"]:
            pages = {}
            for _, h in study_hits[study_hits["term_type"] == term_type].iterrows():
                pages.setdefault(h["term"], []).append(str(h["page"]))
            found[term_type] = "; ".join(f"{t} (p. {', '.join(dict.fromkeys(p))})" for t, p in pages.items())
        first = study_hits.iloc[0]
        ws.append([pdf_files[study_id][:3], f"{first['first_author']} {first['year']}", study_id, pdf_files[study_id],
                   found["name"], found["generic"], None, None, None, None])

    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].hyperlink = row[0].value  # relative link, works because the sheet sits next to the PDFs

    dv = DataValidation(type="list", formula1='"yes,no,unclear"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G2:I{ws.max_row}")

    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", bold=cell.row == 1,
                             color="0563C1" if cell.hyperlink else None, underline="single" if cell.hyperlink else None)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=10):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
    for col, width in zip("ABCDEFGHIJ", [6, 22, 16, 40, 40, 40, 20, 20, 14, 40]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "C2"

    info = wb.create_sheet("Instructions")
    for line in [
        "Fill in the yellow cells only (columns G–J of the Coding sheet); choose yes / no / unclear from the dropdown.",
        "Reviewers 1 and 2 code independently; disagreements are resolved by discussion and entered under Consensus.",
        "",
        "Evaluates AI companion application = yes if the article evaluates a product designed primarily for companionship, "
        "social connection, or role-play rather than clinical care (e.g., Replika, Character.AI, Chai, XiaoIce), "
        "including user-created characters on such platforms, even if they are presented as therapists.",
        "Mentions in the background, related work, discussion or reference list only = no.",
        "",
        "Example (fictitious): Doe 2024 | Reviewer 1: no | Reviewer 2: no | Consensus: no | Notes: Replika only cited in related work (p. 2) and references",
        "",
        "Highlights in the PDFs: yellow = application name, green = generic term. "
        "In Adobe Acrobat Reader, open the Comment panel to jump between highlights.",
    ]:
        info.append([line])
    info.column_dimensions["A"].width = 140
    for row in info.iter_rows():
        row[0].font = Font(name="Arial")
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(path)


def main():
    """Main function to run the full-text search for companion applications."""
    parser = argparse.ArgumentParser(description="Full-text search for companion applications (Supplementary Notes S5.2.3)")
    parser.add_argument("--input", default="data/processed/final_data.csv",
                        help="Path to final_data.csv (default: data/processed/final_data.csv)")
    parser.add_argument("--terms", default="data/auxiliary/companion_screen_terms.yaml",
                        help="Path to the search terms (default: data/auxiliary/companion_screen_terms.yaml)")
    parser.add_argument("--pdf-dir", required=True, help="Folder with the full-text PDFs (not distributed)")
    parser.add_argument("--pdf-index", required=True,
                        help="CSV mapping study_id to first_author, year and pdf_filename (rows in search order)")
    parser.add_argument("--output", default="data/intermediate/companion_screen_hits.csv",
                        help="Path of the hits CSV (default: data/intermediate/companion_screen_hits.csv)")
    parser.add_argument("--highlight-dir", default=None,
                        help="Optional folder for highlighted PDFs and the reviewers' coding sheet")
    args = parser.parse_args()

    df = pd.read_csv(args.input, dtype=str, keep_default_na=False)
    client_facing = set(df.loc[df["application_type"] == "Client-facing application", "study_id"])
    index = pd.read_csv(args.pdf_index, dtype=str, keep_default_na=False)
    index = index[index["study_id"].isin(client_facing)]
    if set(index["study_id"]) != client_facing:
        sys.exit(f"Error: PDF index is missing articles: {sorted(client_facing - set(index['study_id']))}")

    terms = load_terms(Path(args.terms))
    highlight_dir = Path(args.highlight_dir) if args.highlight_dir else None
    if highlight_dir:
        highlight_dir.mkdir(parents=True, exist_ok=True)
        for f in highlight_dir.glob("*.pdf"):  # highlighted PDFs are replaced on every run
            f.unlink()

    hits = []
    for study in index.to_dict("records"):
        highlight_path = highlight_dir / study["pdf_filename"] if highlight_dir else None
        hits += search_pdf(Path(args.pdf_dir) / study["pdf_filename"], study, terms, highlight_path)
    hits = pd.DataFrame(hits)

    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    hits.to_csv(args.output, index=False)

    if highlight_dir:
        coding_path = highlight_dir / CODING_SHEET
        if coding_path.exists():
            print(f"{coding_path} already exists and was not overwritten (delete it to create a new one).")
        else:
            create_coding_sheet(hits, dict(zip(index["study_id"], index["pdf_filename"])), coding_path)
            print(f"Coding sheet saved to {coding_path}")

    print(f"Articles searched: {len(index)}")
    print(f"Hits: {len(hits)} in {hits['study_id'].nunique()} articles")
    print(f"Hits saved to {args.output}")


if __name__ == "__main__":
    main()
